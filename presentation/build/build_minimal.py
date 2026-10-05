"""Builds memory-kit-minimal.pptx: 1 live slide (the capture/curate/recall loop) + 1 hidden reference slide.

Requires Windows with desktop PowerPoint (SVG icons are embedded through COM), Node.js/npm
(Lucide icons, ISC license) and `pip install python-pptx pywin32`.
Usage: python build_minimal.py [output.pptx]   (default: ../memory-kit-minimal.pptx)
Close the output file in PowerPoint first, otherwise saving fails with PermissionError.
"""
import os
import shutil
import subprocess
import sys
import tempfile

import win32com.client
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "memory-kit-minimal.pptx"))
CACHE = os.path.join(tempfile.gettempdir(), "memkit_deck")
ICON_PKG = "lucide-static@1.52.0"  # pinned: icon names are looked up by file name
ICON_SRC = os.path.join(CACHE, "icons", "node_modules", "lucide-static", "icons")
SVG_DIR = os.path.join(CACHE, "svg")
PNG_DIR = os.path.join(CACHE, "png_min")  # slide previews for visual checks

NAVY, INK, MUTED, SOFT, WHITE, SLATE = "0F172A", "1E293B", "64748B", "F1F5F9", "FFFFFF", "475569"
L0, L1, L2, L3 = "E4572E", "F3A712", "2A9D8F", "1D4E89"
L1_TXT = "B7791F"
FONT, MONO = "Segoe UI", "Consolas"
W, H = 13.333, 7.5
M = 0.5
CW = W - 2 * M
TOP, MID = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE
CENTER, RIGHT = PP_ALIGN.CENTER, PP_ALIGN.RIGHT

ICONS = []


def rgb(h):
    return RGBColor.from_string(h)


def _bullet(p, color, indent):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def write(tf, paras):
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for r in spec.get("runs") or [{"text": spec["text"]}]:
            def g(key, default):
                return r.get(key, spec.get(key, default))
            run = p.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(g("size", 14))
            f.bold = g("bold", False)
            f.italic = g("italic", False)
            f.color.rgb = rgb(g("color", INK))
            f.name = g("font", FONT)
            if g("spc", None):
                f._rPr.set("spc", str(g("spc", 0)))
        if "align" in spec:
            p.alignment = spec["align"]
        if "before" in spec:
            p.space_before = Pt(spec["before"])
        if "after" in spec:
            p.space_after = Pt(spec["after"])
        if "bullet" in spec:
            _bullet(p, spec["bullet"], spec.get("indent", 0.18))
        if "hang" in spec:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(spec["hang"]))))
            pPr.set("indent", str(-int(Inches(spec["hang"]))))


def textbox(slide, x, y, w, h, paras, anchor=TOP, pad=0.04):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    write(tf, paras)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE, radius=None, line_w=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if radius is not None:
        s.adjustments[0] = radius
    return s


def circle(slide, cx, cy, d, fill=None, line=None, line_w=2.0):
    return rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, kind=MSO_SHAPE.OVAL, line_w=line_w)


def icon(slide_no, name, color, cx, cy, size):
    ICONS.append(dict(slide=slide_no, name=name, color=color, cx=cx, cy=cy, size=size))


def gradient(shape, colors):
    shape.fill.gradient()
    shape.fill.gradient_angle = 0
    gs_lst = shape._element.spPr.find(qn("a:gradFill")).find(qn("a:gsLst"))
    for gs in list(gs_lst):
        gs_lst.remove(gs)
    for i, c in enumerate(colors):
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(int(100000 * i / (len(colors) - 1))))
        etree.SubElement(gs, qn("a:srgbClr")).set("val", c)


def _arrow_line(shape, color, width, arrow=True):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    ln = shape._element.spPr.find(qn("a:ln"))
    ln.set("cap", "rnd")
    if arrow:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def straight(slide, x1, y1, x2, y2, color, width=2.75, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _arrow_line(c, color, width, arrow)


def polyline(slide, pts, color, width=2.75):
    emu = [(int(Inches(x)), int(Inches(y))) for x, y in pts]
    fb = slide.shapes.build_freeform(*emu[0], scale=1.0)
    fb.add_line_segments(emu[1:], close=False)
    shp = fb.convert_to_shape()
    shp.shadow.inherit = False
    shp.fill.background()
    _arrow_line(shp, color, width)


def station(slide, slide_no, cx, cy, color, name, label, above=False):
    circle(slide, cx, cy, 0.74, fill=WHITE, line=color, line_w=2.5)
    icon(slide_no, name, color, cx, cy, 0.4)
    ly = cy - 0.76 if above else cy + 0.42
    textbox(slide, cx - 0.9, ly, 1.8, 0.34, [dict(text=label, size=15, bold=True, color=color, align=CENTER)])


def top_strip(slide):
    seg = W / 4
    for i, c in enumerate((L0, L1, L2, L3)):
        rect(slide, i * seg, 0, seg + 0.01, 0.09, fill=c)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "memory-kit — Knowledge management when coding with agents"
prs.core_properties.author = ""
prs.core_properties.last_modified_by = ""
blank = prs.slide_layouts[6]

# ================================================================ slide 1: live slide
s1 = prs.slides.add_slide(blank)
top_strip(s1)
textbox(s1, M, 0.35, 6, 0.3, [dict(text="memory-kit", size=12, bold=True, color=L3, spc=150)])
textbox(s1, M, 0.6, 9, 0.8, [dict(text="Chat forgets. Memory doesn't.", size=34, bold=True, color=NAVY)])

chip_x, chip_y, chip_w, chip_h = W - M - 2.0, 0.66, 2.0, 0.52
rect(s1, chip_x, chip_y, chip_w, chip_h, fill=SOFT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
icon(1, "file-text", NAVY, chip_x + 0.38, chip_y + chip_h / 2, 0.3)
textbox(s1, chip_x + 0.6, chip_y, chip_w - 0.64, chip_h,
        [dict(text="MEMORY.md", size=14, bold=True, font=MONO, color=NAVY)], anchor=MID)

CY, D = 3.75, 1.3
R = D / 2
CHAT_X = 1.55
XS = (5.0, 7.3, 9.6, 11.9)
BADGE = R * 0.7071
MID_X = (XS[0] + XS[3]) / 2
CUR_Y, REC_Y = 2.2, 5.7  # Curate bus above all levels, Recall bus below
LBL_BOTTOM = CY + R + 0.06 + 0.62
CHAT_END = CY + R + 0.06 + 0.4 + 0.08

bar = rect(s1, XS[0], CY - 0.08, XS[3] - XS[0], 0.16)
gradient(bar, [L0, L1, L2, L3])
straight(s1, CHAT_X + R + 0.08, CY, XS[0] - R - 0.08, CY, L0)

# Curate: requested in the chat, then acts on every level
straight(s1, CHAT_X, CY - R - 0.08, CHAT_X, CUR_Y + 0.37 + 0.08, L2)
straight(s1, CHAT_X, CUR_Y, XS[3], CUR_Y, L2, arrow=False)
for x in XS:
    straight(s1, x, CUR_Y, x, CY - R - 0.08, L2)

# Recall gathers from every level, then feeds the chat
for x in XS:
    straight(s1, x, LBL_BOTTOM + 0.08, x, REC_Y, L3, arrow=False)
polyline(s1, [(XS[3], REC_Y), (CHAT_X, REC_Y), (CHAT_X, CHAT_END)], L3)

circle(s1, CHAT_X, CY, D, fill=SLATE)
icon(1, "messages-square", WHITE, CHAT_X, CY, 0.62)
textbox(s1, CHAT_X - 0.8, CY + R + 0.06, 1.6, 0.4, [dict(text="Chat", size=16, bold=True, color=SLATE, align=CENTER)])

for x, (lvl, name, fill, txt, fg, ic) in zip(XS, (
    ("L0", "Scratchpad", L0, L0, WHITE, "notebook-pen"),
    ("L1", "Task", L1, L1_TXT, NAVY, "list-checks"),
    ("L2", "Module", L2, L2, WHITE, "boxes"),
    ("L3", "Global", L3, L3, WHITE, "book-open"),
)):
    circle(s1, x, CY, D, fill=fill)
    icon(1, ic, fg, x, CY, 0.62)
    textbox(s1, x - 0.9, CY + R + 0.06, 1.8, 0.62, [
        dict(text=lvl, size=16, bold=True, color=txt, align=CENTER),
        dict(text=name, size=12, color=MUTED, align=CENTER),
    ])

for x, dx, color, ic in ((XS[0], -BADGE, L0, "flame"), (XS[3], BADGE, L3, "snowflake")):
    circle(s1, x + dx, CY - BADGE, 0.46, fill=WHITE, line=color, line_w=2)
    icon(1, ic, color, x + dx, CY - BADGE, 0.27)

station(s1, 1, (CHAT_X + R + XS[0] - R) / 2, CY, L0, "camera", "Capture")
station(s1, 1, CHAT_X, CUR_Y, L2, "sparkles", "Curate", above=True)
circle(s1, CHAT_X - 0.36, CUR_Y + 0.32, 0.38, fill=L2, line=WHITE, line_w=1.5)
icon(1, "user-check", WHITE, CHAT_X - 0.36, CUR_Y + 0.32, 0.22)
station(s1, 1, MID_X, REC_Y, L3, "search", "Recall")

textbox(s1, M, 6.62, CW, 0.5, [dict(text="Why review code when you can validate knowledge?", size=20, italic=True,
                                     color=NAVY, align=CENTER)], anchor=MID)

s1.notes_slide.notes_text_frame.text = """\
TALK TRACK — 15 min. The slide is one loop: walk it left → right → back along the bottom.

[0:00–2:00] HOOK — grey "Chat" circle
- Who has lost a good agent session to a context reset, or to a compaction that kept the wrong things?
- Chat grows fast and gets noisy while iterating. Compaction shrinks it, but doesn't keep only what matters.
- Agents rush to write new code instead of first gathering what we already know.

[2:00–3:30] THE IDEA — MEMORY.md chip
- Move knowledge out of the chat into a structured memory, and force agents to use it.
- The whole system is one Markdown file: the model, the rules, three verbs.
- Works with any agent (Copilot, Claude Code, Cursor, Aider). No build, no dependencies.

[3:30–5:30] CAPTURE — red arrow (camera)
- Before a reset or a handoff: append a raw checkpoint to L0 — state, actions, what worked, what failed, hypotheses, next steps.
- "Capture fast, curate later": no analysis, so it is cheap enough to actually happen.

[5:30–8:30] FOUR LEVELS — flame (hot) to snowflake (cold)
- L0 Scratchpad: raw, git-ignored, ephemeral.
- L1 Task: working signal for the current task, cleaned when the task ends.
- L2 Module: durable lessons beside the code, including known failures.
- L3 Global: curated, cross-cutting truth in docs/ — can be HTML for humans.
- Authority: colder wins — L3 > L2 > L1 > L0 > Chat. If docs and code disagree, trust the running code, then fix the docs.

[8:30–11:00] CURATE — teal line leaving the chat (sparkles + approval badge), reaching all four levels
- You ask for it in the chat; Curate then looks at every level at once: L0 to L3.
- Promotion must be earned: validated, reusable, stable, not contradicted.
- Cleanup over accumulation: keep, merge, summarize, promote, archive, delete, reset.
- Failures are valuable: anti-patterns move up as known failures — never silently deleted.
- Approval-first: the agent proposes a plan per file and stops. You decide what becomes durable.

[11:00–13:00] RECALL — blue line collecting from all four levels, back to the chat (magnifier)
- At the start of every task: a read-only briefing built from L0 to L3 plus the code, delivered into the chat.
- Separates facts from uncertainty, lists failed paths, surfaces conflicts, declares missing sources.
- This is the "think first" step agents skip. And token economy: a fresh session reads a few curated files instead of a bloated chat.

[13:00–15:00] CLOSE — the question at the bottom
- Adoption in ~1 minute: copy the folder, add one line to your agent instructions, set two paths.
- When agents write most of the code, review shifts to knowledge: decisions, constraints, known failures.
- "Why review code when you can validate knowledge?" → Q&A. Details on hidden slide 2 (type 2 + Enter)."""

# ================================================================ slide 2: hidden reference
s2 = prs.slides.add_slide(blank)
s2._element.set("show", "0")
top_strip(s2)
textbox(s2, M, 0.3, 8, 0.3, [dict(text="BACKUP  ·  HIDDEN IN SLIDESHOW", size=11, bold=True, color=L3, spc=150)])
textbox(s2, M, 0.52, 8, 0.6, [dict(text="memory-kit — the details", size=24, bold=True, color=NAVY)])
textbox(s2, W - M - 6.0, 0.6, 6.0, 0.45, [dict(runs=[
    dict(text="MEMORY.md", font=MONO, bold=True, color=NAVY),
    dict(text="  ·  single source of truth  ·  no build  ·  any agent", color=MUTED),
], size=11, align=RIGHT)], anchor=MID)

GAP = 0.18
cw4 = (CW - 3 * GAP) / 4
cw3 = (CW - 2 * GAP) / 3

ay, ah = 1.25, 1.6
for i, (lvl, name, temp, fill, txt, fg, ic, path, where, life, holds) in enumerate((
    ("L0", "Scratchpad", "hot · raw", L0, L0, WHITE, "notebook-pen", "notes/L0_*.md", "  git-ignored",
     "Ephemeral — triage or delete", "Raw tries, errors, hypotheses"),
    ("L1", "Task memory", "warm · working", L1, L1_TXT, NAVY, "list-checks", "notes/L1_*.md", "",
     "Per task — cleaned at the end", "Working signal for the task"),
    ("L2", "Module memory", "cool · stable", L2, L2, WHITE, "boxes", "L2_*.md", "  beside the code",
     "Durable · module-scoped", "Lessons & known failures"),
    ("L3", "Global memory", "cold · curated", L3, L3, WHITE, "book-open", "docs/", "  Markdown or HTML",
     "Durable · cross-cutting", "Architecture & decisions"),
)):
    x = M + i * (cw4 + GAP)
    rect(s2, x, ay, cw4, ah, fill=SOFT)
    circle(s2, x + 0.42, ay + 0.4, 0.52, fill=fill)
    icon(2, ic, fg, x + 0.42, ay + 0.4, 0.28)
    textbox(s2, x + 0.78, ay + 0.12, cw4 - 0.85, 0.58, [
        dict(runs=[dict(text=lvl + "  ", color=txt), dict(text=name, color=NAVY)], size=13, bold=True),
        dict(text=temp, size=10, color=MUTED),
    ])
    for j, (ric, runs) in enumerate((
        ("folder-open", [dict(text=path, font=MONO, bold=True, size=10, color=NAVY), dict(text=where, color=MUTED)]),
        ("hourglass", [dict(text=life)]),
        ("file-text", [dict(text=holds)]),
    )):
        ry = ay + 0.78 + j * 0.26
        icon(2, ric, MUTED, x + 0.3, ry + 0.13, 0.18)
        textbox(s2, x + 0.46, ry, cw4 - 0.52, 0.26, [dict(runs=runs, size=10.5, color=INK)], anchor=MID)

by, bh = ay + ah + 0.15, 2.3
for i, (verb, ic, c, writes, quote, steps, cmd) in enumerate((
    ("RECALL", "search", L3, "read-only", "“What do we already know about X?”", (
        "Inventory chat → L0 → L1 → L2 → L3 → instructions → code",
        "State unavailable sources and lower confidence",
        "Briefing: facts vs. uncertainty, failed paths, conflicts & drift, next step",
    ), "/mem-recall <topic>"),
    ("CAPTURE", "camera", L0, "appends L0", "“Save the context before we reset.”", (
        "Append — never overwrite — to notes/L0_<topic>.md",
        "Checkpoint: state · actions · observations · hypotheses · next steps · raw snippets",
        "No analysis · keep contradictions",
    ), "/mem-capture <topic>"),
    ("CURATE", "sparkles", L2, "after approval", "“Clean up and promote the useful notes.”", (
        "Detect drift, duplication, stale notes, noise, missing anti-patterns",
        "Plan per file: level · action · reason · benefit · risk",
        "STOP for approval → run the approved subset → report",
    ), "/mem-curate <scope>"),
)):
    x = M + i * (cw3 + GAP)
    rect(s2, x, by, cw3, bh, fill=SOFT)
    circle(s2, x + 0.42, by + 0.4, 0.52, fill=c)
    icon(2, ic, WHITE, x + 0.42, by + 0.4, 0.28)
    textbox(s2, x + 0.78, by + 0.12, cw3 - 0.9, 0.58, [
        dict(text=verb, size=13, bold=True, color=c, spc=80),
        dict(text=quote, size=10, italic=True, color=MUTED),
    ])
    textbox(s2, x + 0.2, by + 0.76, cw3 - 0.4, bh - 1.15,
            [dict(text=t, size=10.5, bullet=c, after=3) for t in steps])
    textbox(s2, x + 0.2, by + bh - 0.38, 2.2, 0.28, [dict(text=cmd, size=10, font=MONO, color=MUTED)], anchor=MID)
    chip = rect(s2, x + cw3 - 1.42, by + bh - 0.38, 1.24, 0.28, fill=WHITE, line=c, line_w=1.25,
                kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = chip.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MID
    write(tf, [dict(text=writes, size=9.5, bold=True, color=c, align=CENTER)])

cy, ch = by + bh + 0.15, 1.45
code = dict(font=MONO, size=9.5, color=NAVY)
for i, (ic, c, title, paras) in enumerate((
    ("trending-up", L2, "Promote only when", [dict(text=t, size=10, bullet=L2) for t in (
        "validated by usage or code", "reusable beyond one case", "stable across resets",
        "not contradicted from above")] + [
        dict(text="task→L1 · module→L2 · project→L3", size=9.5, font=MONO, color=MUTED, before=3)]),
    ("scale", L3, "Authority on conflict", [
        dict(runs=[dict(text="L3", color=L3), dict(text=" > ", color=MUTED), dict(text="L2", color=L2),
                   dict(text=" > ", color=MUTED), dict(text="L1", color=L1_TXT), dict(text=" > ", color=MUTED),
                   dict(text="L0", color=L0), dict(text=" > ", color=MUTED), dict(text="Chat", color=SLATE)],
             size=13, bold=True, after=3),
        dict(text="Chat is never authoritative", size=10, bullet=L3),
        dict(text="Code beats docs — then fix the docs", size=10, bullet=L3),
        dict(text="Lower never overwrites higher", size=10, bullet=L3),
    ]),
    ("trash-2", L0, "Cleanup over accumulation", [
        dict(text="keep · merge · summarize · promote", size=9, font=MONO, bold=True, color=L0),
        dict(text="archive · delete · reset", size=9, font=MONO, bold=True, color=L0, after=4),
        dict(text="Never silently delete anti-patterns — relocate them to L2/L3", size=10),
    ]),
    ("rocket", NAVY, "Install in ~1 minute", [
        dict(runs=[dict(text="1\t", bold=True), dict(text="Copy "), dict(text="memory-kit/", **code),
                   dict(text=" into the repo")], size=10, hang=0.2),
        dict(runs=[dict(text="2\t", bold=True), dict(text="Point "), dict(text="AGENTS.md", **code),
                   dict(text=" / "), dict(text="CLAUDE.md", **code), dict(text=" to "),
                   dict(text="MEMORY.md", **code)], size=10, before=2, hang=0.2),
        dict(runs=[dict(text="3\t", bold=True), dict(text="Set 2 dirs; git-ignore "),
                   dict(text="L0_*", **code)], size=10, before=2, hang=0.2),
    ]),
)):
    x = M + i * (cw4 + GAP)
    rect(s2, x, cy, cw4, ch, fill=SOFT)
    icon(2, ic, c, x + 0.32, cy + 0.25, 0.24)
    textbox(s2, x + 0.5, cy + 0.08, cw4 - 0.6, 0.34, [dict(text=title, size=12, bold=True, color=NAVY)], anchor=MID)
    textbox(s2, x + 0.2, cy + 0.45, cw4 - 0.4, ch - 0.5, paras)

textbox(s2, M, 7.05, 9, 0.3, [dict(text="memory-kit  ·  reference for Q&A", size=10, color=MUTED)])

s2.notes_slide.notes_text_frame.text = """\
HIDDEN BACKUP SLIDE — skipped during the slideshow.
- Q&A: during the slideshow, type 2 then Enter to jump here; type 1 then Enter to go back.
- Also usable as a one-page handout (File > Print includes hidden slides when "Print hidden slides" is checked)."""


# ================================================================ embed SVG icons via PowerPoint, render previews
def ensure_icons():
    if os.path.isdir(ICON_SRC):
        return
    npm = shutil.which("npm")
    if not npm:
        sys.exit("npm not found: install Node.js to fetch the Lucide icons")
    subprocess.run([npm, "install", ICON_PKG, "--prefix", os.path.join(CACHE, "icons"), "--no-audit", "--no-fund"],
                   check=True)


def colored_svg(name, color):
    os.makedirs(SVG_DIR, exist_ok=True)
    dst = os.path.join(SVG_DIR, f"{name}-{color}.svg")
    with open(os.path.join(ICON_SRC, f"{name}.svg"), encoding="utf-8") as f:
        svg = f.read().replace('stroke="currentColor"', f'stroke="#{color}"')
    with open(dst, "w", encoding="utf-8") as f:
        f.write(svg)
    return dst


ensure_icons()
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)

app = win32com.client.DispatchEx("PowerPoint.Application")
pres = app.Presentations.Open(OUT, 0, 0, 0)
try:
    for ic in ICONS:
        size = ic["size"] * 72
        pic = pres.Slides.Item(ic["slide"]).Shapes.AddPicture(
            colored_svg(ic["name"], ic["color"]), 0, -1,
            (ic["cx"] - ic["size"] / 2) * 72, (ic["cy"] - ic["size"] / 2) * 72, size, size)
        pic.Name = f"icon {ic['name']}"
        pic.AlternativeText = f"{ic['name'].replace('-', ' ')} icon"
    pres.Save()
    os.makedirs(PNG_DIR, exist_ok=True)
    for i in range(1, pres.Slides.Count + 1):
        pres.Slides.Item(i).Export(os.path.join(PNG_DIR, f"slide{i}.png"), "PNG", 1600, 900)
finally:
    pres.Close()
    if app.Presentations.Count == 0:
        app.Quit()

print(OUT, f"({len(ICONS)} icons)")
