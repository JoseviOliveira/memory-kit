"""Builds memory-kit-knowledge-management.pptx: 3 slides (executive summary, model, workflow) with timed notes.

Requires `pip install python-pptx`; no PowerPoint needed.
Usage: python build_overview.py [output.pptx]   (default: ../memory-kit-knowledge-management.pptx)
Close the output file in PowerPoint first, otherwise saving fails with PermissionError.
"""
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(
    sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "memory-kit-knowledge-management.pptx"))
OUT_DIR = os.path.dirname(OUT)

NAVY, INK, MUTED, SOFT, LINE, WHITE, GREY = "0F172A", "1E293B", "64748B", "F1F5F9", "CBD5E1", "FFFFFF", "94A3B8"
L0, L1, L2, L3 = "E4572E", "F3A712", "2A9D8F", "1D4E89"
L1_TXT = "B7791F"
FONT, MONO = "Segoe UI", "Consolas"
W, H = 13.333, 7.5
M = 0.5
CW = W - 2 * M
TOP, MID = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE
CENTER, RIGHT = PP_ALIGN.CENTER, PP_ALIGN.RIGHT


def rgb(h):
    return RGBColor.from_string(h)


def _bullet(p, color, indent):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def write(tf, paras):
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for r in spec.get("runs") or [{"text": spec["text"]}]:
            def g(key, default):
                return r.get(key, spec.get(key, default))
            run = p.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(g("size", 14))
            f.bold = g("bold", False)
            f.italic = g("italic", False)
            f.color.rgb = rgb(g("color", INK))
            f.name = g("font", FONT)
            if g("spc", None):
                f._rPr.set("spc", str(g("spc", 0)))
        if "align" in spec:
            p.alignment = spec["align"]
        if "before" in spec:
            p.space_before = Pt(spec["before"])
        if "after" in spec:
            p.space_after = Pt(spec["after"])
        if "bullet" in spec:
            _bullet(p, spec["bullet"], spec.get("indent", 0.2))


def textbox(slide, x, y, w, h, paras, anchor=TOP, pad=0.04):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    write(tf, paras)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE, radius=None, line_w=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if radius is not None:
        s.adjustments[0] = radius
    return s


def label_in(shape, paras):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MID
    tf.word_wrap = False
    write(tf, paras)


def gradient(shape, colors):
    shape.fill.gradient()
    shape.fill.gradient_angle = 0
    gs_lst = shape._element.spPr.find(qn("a:gradFill")).find(qn("a:gsLst"))
    for gs in list(gs_lst):
        gs_lst.remove(gs)
    for i, c in enumerate(colors):
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(int(100000 * i / (len(colors) - 1))))
        etree.SubElement(gs, qn("a:srgbClr")).set("val", c)


def frame(slide, num, kicker, title, notes):
    seg = W / 4
    for i, c in enumerate((L0, L1, L2, L3)):
        rect(slide, i * seg, 0, seg + 0.01, 0.09, fill=c)
    textbox(slide, M, 0.32, CW, 0.3, [dict(text=kicker, size=12, bold=True, color=L3, spc=150)])
    textbox(slide, M, 0.6, CW, 0.7, [dict(text=title, size=30, bold=True, color=NAVY)])
    rect(slide, M, 7.0, CW, 0.012, fill=LINE)
    textbox(slide, M, 7.05, 9, 0.3,
            [dict(text="memory-kit  ·  Knowledge management when coding with agents", size=10, color=MUTED)])
    textbox(slide, W - M - 1.5, 7.05, 1.5, 0.3, [dict(text=f"{num} / 3", size=10, color=MUTED, align=RIGHT)])
    slide.notes_slide.notes_text_frame.text = notes


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "memory-kit — Knowledge management when coding with agents"
prs.core_properties.author = ""
prs.core_properties.last_modified_by = ""
blank = prs.slide_layouts[6]

# ---------------------------------------------------------------- slide 1: executive summary
s = prs.slides.add_slide(blank)
frame(s, 1, "EXECUTIVE SUMMARY", "memory-kit — durable knowledge for AI-assisted coding", """\
[~4 min · 0:00 – 4:00]

OPENING HOOK
Quick show of hands: who has lost a productive agent session to a context reset — or to a compaction that kept the wrong things?

ORIGIN
memory-kit came out of building a web app from scratch with Codex and Claude agents. Two observations:
1. Agents are eager to produce new code, but weak at thinking first — gathering what is already known before deciding.
2. The chat window is a poor place to keep knowledge: it grows fast, gets dirty while iterating, and compaction shrinks it without guaranteeing it keeps only what matters.
The question became: how do we move knowledge out of the chat into a durable, organized memory — and force agents to use it?

THE FOUR NUMBERS
- 1 file: MEMORY.md defines the model, the rules and the verbs. One source of truth, no drift.
- 4 levels: from hot/raw scratch notes to cold/curated docs.
- 3 verbs: capture, recall, curate.
- 0 dependencies: plain Markdown, works with Copilot, Claude Code, Cursor, Aider — about one minute to install.

VALUE
Less noise and fewer tokens; dead ends are preserved so we do not repeat them; and the memory becomes the interface between humans and agents.

Leave the provocative thought on screen — we come back to it at the end.""")

tiles = [
    ("1", "file", "MEMORY.md holds the model, rules and verbs", L0, L0),
    ("4", "memory levels", "L0 → L3, from hot/raw to cold/curated", L1, L1_TXT),
    ("3", "verbs", "capture · recall · curate", L2, L2),
    ("0", "dependencies", "any agent · ~1-minute install", L3, L3),
]
tw, ty, th = 2.9, 1.55, 1.05
gap = (CW - 4 * tw) / 3
for i, (n, label, sub, bar, txt) in enumerate(tiles):
    x = M + i * (tw + gap)
    rect(s, x, ty, tw, th, fill=SOFT)
    rect(s, x, ty, 0.08, th, fill=bar)
    textbox(s, x + 0.2, ty, 0.75, th, [dict(text=n, size=40, bold=True, color=txt)], anchor=MID)
    textbox(s, x + 0.95, ty, tw - 1.05, th,
            [dict(text=label, size=15, bold=True, color=NAVY), dict(text=sub, size=11, color=MUTED, before=2)],
            anchor=MID)

cols = [
    ("The problem", L0, [
        "Chat sessions grow fast and get noisy while iterating",
        "Compaction shrinks context — it doesn't keep only what matters",
        "Agents rush to write code instead of gathering knowledge first",
        "Every reset loses decisions, dead ends — and tokens",
    ]),
    ("The approach", L3, [
        "A 4-level memory: scratchpad → task → module → global docs",
        "Knowledge is promoted only when validated, reusable and stable",
        "Explicit authority: curated docs and running code beat chat",
        "Plain Markdown — works with Copilot, Claude Code, Cursor, Aider…",
    ]),
    ("The value", L2, [
        "Less noise, fewer tokens, faster restarts after a reset",
        "Failed approaches are kept, so mistakes aren't repeated",
        "A shared, reviewable interface between humans and agents",
        "Teams validate knowledge, not only code",
    ]),
]
cw, cy, ch = 3.95, 2.85, 3.15
gap = (CW - 3 * cw) / 2
for i, (title, c, items) in enumerate(cols):
    x = M + i * (cw + gap)
    rect(s, x, cy, cw, ch, fill=WHITE, line=LINE)
    rect(s, x, cy, cw, 0.07, fill=c)
    textbox(s, x + 0.2, cy + 0.2, cw - 0.4, 0.45, [dict(text=title, size=18, bold=True, color=c)])
    textbox(s, x + 0.2, cy + 0.72, cw - 0.4, ch - 0.85,
            [dict(text=t, size=13, color=INK, bullet=c, after=8) for t in items])

rect(s, M, 6.18, CW, 0.62, fill=NAVY, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
textbox(s, M + 0.3, 6.18, CW - 0.6, 0.62, [dict(runs=[
    dict(text="PROVOCATIVE THOUGHT     ", size=11, bold=True, color=L1, spc=150),
    dict(text="“Why spend time reviewing code when you can validate knowledge directly from the memory kit?”",
         size=14, italic=True, color=WHITE),
])], anchor=MID)

# ---------------------------------------------------------------- slide 2: the model
s = prs.slides.add_slide(blank)
frame(s, 2, "THE MODEL  ·  L0 → L3", "Knowledge flows from hot to cold", """\
[~5.5 min · 4:00 – 9:30]

THE ARROW
Read it left to right: knowledge starts hot and raw, and cools down as it gets validated.

THE FOUR LEVELS
- L0 Scratchpad: raw dumps, git-ignored, cheap to write. "Capture fast, curate later."
- L1 Task memory: what is still useful for the current task; cleaned when the task ends.
- L2 Module memory: lives beside the code; durable lessons for one module, including known failures.
- L3 Global memory: curated, cross-cutting docs. Can be HTML — this is the level humans read.

ILLUSTRATIVE EXAMPLE (token-refresh bug)
L0 gets the raw attempts and stack traces. L1 keeps "approach B works, A fails because of X". L2_auth.md keeps "known failure: don't do A". The L3 architecture doc records the final token-rotation design.

THREE RULES MAKE IT WORK
1. Promotion must be earned: validated, reusable, stable, not contradicted. Speculation stays hot.
2. Authority is explicit: L3 > L2 > L1 > L0 > Chat. Chat is never authoritative until validated against files. If docs and code disagree, trust the running code, then propose a docs fix.
3. Cleanup over accumulation: stale, duplicated notes are liabilities. But failures are valuable — anti-patterns are relocated as known failures, never silently deleted.

TOKEN ANGLE
A fresh session reads a few small, curated files instead of dragging a bloated chat history along.""")

arrow = rect(s, M, 1.5, CW, 0.55, kind=MSO_SHAPE.RIGHT_ARROW)
arrow.adjustments[0] = 0.72
arrow.adjustments[1] = 0.6
gradient(arrow, [L0, L1, L2, L3])
textbox(s, M + 0.2, 1.5, 5, 0.55, [dict(text="HOT  ·  raw, cheap to write", size=12, bold=True, color=WHITE, spc=80)],
        anchor=MID)
textbox(s, W - M - 0.5 - 5, 1.5, 5, 0.55,
        [dict(text="COLD  ·  curated, trusted", size=12, bold=True, color=WHITE, spc=80, align=RIGHT)], anchor=MID)

levels = [
    ("L0", "Scratchpad", "hot · raw", L0, WHITE, "notes/L0_*.md", "  (git-ignored)",
     "Ephemeral — triage soon or delete", "Raw checkpoints: what worked, what failed, errors"),
    ("L1", "Task memory", "warm · working", L1, NAVY, "notes/L1_*.md", "",
     "Per task — cleaned at task end", "Working signal and decisions for the current task"),
    ("L2", "Module memory", "cool · stable", L2, WHITE, "L2_*.md", "  beside the code",
     "Durable, module-scoped", "Module lessons and known failures (anti-patterns)"),
    ("L3", "Global memory", "cold · curated", L3, WHITE, "docs/", "  (Markdown or HTML)",
     "Durable, cross-cutting", "Project truths: architecture, conventions, decisions"),
]
kw, ky, kh = 2.9, 2.22, 2.93
gap = (CW - 4 * kw) / 3
for i, (lvl, name, temp, c, fg, mono, suffix, life, holds) in enumerate(levels):
    x = M + i * (kw + gap)
    rect(s, x, ky, kw, kh, fill=WHITE, line=LINE)
    rect(s, x, ky, kw, 0.8, fill=c)
    textbox(s, x + 0.15, ky, 0.75, 0.8, [dict(text=lvl, size=28, bold=True, color=fg)], anchor=MID)
    textbox(s, x + 0.85, ky, kw - 0.9, 0.8,
            [dict(text=name, size=16, bold=True, color=fg), dict(text=temp, size=11, color=fg)], anchor=MID)
    body = []
    for label, runs in (
        ("WHERE", [dict(text=mono, font=MONO, bold=True, color=NAVY), dict(text=suffix, size=11, color=MUTED)]),
        ("LIFESPAN", [dict(text=life)]),
        ("HOLDS", [dict(text=holds)]),
    ):
        body.append(dict(text=label, size=9.5, bold=True, color=MUTED, spc=100, before=10))
        body.append(dict(runs=runs, size=12.5, color=INK))
    textbox(s, x + 0.15, ky + 0.82, kw - 0.3, kh - 0.9, body)
    if i < 3:
        rect(s, x + kw + gap / 2 - 0.08, ky + 0.25, 0.16, 0.3, fill=GREY, kind=MSO_SHAPE.CHEVRON)

bw, by, bh = 3.95, 5.32, 1.53
gap = (CW - 3 * bw) / 2
xs = [M + i * (bw + gap) for i in range(3)]
for x, c in zip(xs, (L2, L3, L0)):
    rect(s, x, by, bw, bh, fill=SOFT)
    rect(s, x, by, 0.07, bh, fill=c)

textbox(s, xs[0] + 0.2, by + 0.1, bw - 0.4, 0.35, [dict(text="Promotion must be earned", size=14, bold=True, color=NAVY)])
textbox(s, xs[0] + 0.2, by + 0.48, bw - 0.4, bh - 0.52, [dict(text=t, size=11.5, bullet=L2, after=1) for t in (
    "Validated by usage or current code",
    "Reusable beyond one case",
    "Stable enough to survive a reset",
    "Not contradicted by higher authority",
)])

textbox(s, xs[1] + 0.2, by + 0.1, bw - 0.4, 0.35, [dict(text="Authority on conflict", size=14, bold=True, color=NAVY)])
cx = xs[1] + 0.2
for j, (t, c, fg) in enumerate((("L3", L3, WHITE), ("L2", L2, WHITE), ("L1", L1, NAVY), ("L0", L0, WHITE),
                                ("Chat", GREY, WHITE))):
    chip = rect(s, cx, by + 0.5, 0.55, 0.32, fill=c, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    label_in(chip, [dict(text=t, size=12, bold=True, color=fg, align=CENTER)])
    if j < 4:
        textbox(s, cx + 0.55, by + 0.5, 0.2, 0.32, [dict(text=">", size=13, bold=True, color=MUTED, align=CENTER)],
                anchor=MID, pad=0)
    cx += 0.75
textbox(s, xs[1] + 0.2, by + 0.9, bw - 0.4, 0.6, [dict(
    text="Chat is never authoritative until validated. Docs and code disagree? Trust the running code.",
    size=11, color=INK)])

textbox(s, xs[2] + 0.2, by + 0.1, bw - 0.4, 0.35, [dict(text="Cleanup over accumulation", size=14, bold=True, color=NAVY)])
textbox(s, xs[2] + 0.2, by + 0.48, bw - 0.4, bh - 0.52, [
    dict(text="keep · merge · summarize · promote · archive · delete · reset", size=10.5, font=MONO, bold=True,
         color=L0, after=4),
    dict(text="Low-signal notes are liabilities. Anti-patterns are relocated to L2/L3 — never silently deleted.",
         size=11, color=INK),
])

# ---------------------------------------------------------------- slide 3: the workflow
s = prs.slides.add_slide(blank)
frame(s, 3, "THE WORKFLOW", "Three verbs, one loop", """\
[~5.5 min · 9:30 – 15:00]

PLAIN LANGUAGE
The three verbs work with plain sentences in any agent. Slash commands in prompts/ are optional sugar for tools that support prompt files.

1. RECALL — start of a task. Read-only.
The agent inventories chat, L0 to L3 and the code needed to validate claims, states which sources were missing, and returns a briefing that keeps confirmed facts separate from uncertainty and surfaces conflicts instead of flattening them. This is the "think first" step agents usually skip.

2. CAPTURE — before any reset or handoff. Append-only.
A raw checkpoint into L0: current state, recent actions, observations, hypotheses, next steps. Deliberately no analysis — it must be cheap so it actually happens.

3. CURATE — end of a task. Approval-first.
The agent proposes a plan per file — keep, merge, summarize, promote, archive, delete, reset — with reason and risk, then STOPS. The human decides what becomes durable knowledge.

THE LOOP
Each task starts with more trusted knowledge than the previous one.

ADOPTION
Copy the folder, add one line to the always-loaded instructions, set two paths. No build, no runtime. Because the model is defined once, there is no drift between tools.

CLOSE
Back to the provocative thought: when agents write most of the code, review effort shifts to validating knowledge — decisions, constraints, known failures. memory-kit makes that knowledge explicit and reviewable.
Q&A.""")

verbs = [
    ("RECALL", "“What do we already know about X?”", L3, "Start of a task · after a reset",
     "Nothing — read-only", (
         "Reads chat, L0 → L3 and the code",
         "Flags missing sources, lowers confidence",
         "Briefing: facts vs. uncertainty, failed paths, drift, next step",
     ), "/mem-recall <topic>"),
    ("CAPTURE", "“Save the context before we reset.”", L0, "Before a reset · checkpoint · handoff",
     "Appends to L0 only", (
         "Raw checkpoint: state, actions, observations, hypotheses, next steps",
         "No analysis, no cleanup, no promotion",
         "Contradictions kept as-is",
     ), "/mem-capture <topic>"),
    ("CURATE", "“Clean up and promote the useful notes.”", L2, "End of a task · notes piling up",
     "Only after your approval", (
         "Detects drift, duplicates, stale notes",
         "Plans per file: action · reason · risk",
         "STOPS for approval, then executes and reports",
     ), "/mem-curate <scope>"),
]
vw, vy, vh = 3.85, 1.5, 3.35
gap = (CW - 3 * vw) / 2
for i, (verb, quote, c, when, writes, steps, cmd) in enumerate(verbs):
    x = M + i * (vw + gap)
    rect(s, x, vy, vw, vh, fill=WHITE, line=LINE)
    rect(s, x, vy, vw, 0.95, fill=c)
    textbox(s, x + 0.2, vy + 0.05, vw - 0.4, 0.85, [
        dict(runs=[dict(text=f"{i + 1}   ", size=20, color=WHITE), dict(text=verb, size=22, bold=True, color=WHITE, spc=100)]),
        dict(text=quote, size=12, italic=True, color=WHITE),
    ], anchor=MID)
    body = [
        dict(runs=[dict(text="WHEN   ", size=9, bold=True, color=MUTED, spc=100), dict(text=when, size=11.5, color=INK)]),
        dict(runs=[dict(text="WRITES   ", size=9, bold=True, color=MUTED, spc=100),
                   dict(text=writes, size=11.5, bold=True, color=c)], before=3, after=8),
    ] + [dict(text=t, size=11.5, bullet=c, after=4) for t in steps]
    textbox(s, x + 0.2, vy + 1.08, vw - 0.4, vh - 1.5, body)
    textbox(s, x + 0.2, vy + vh - 0.4, vw - 0.4, 0.3, [dict(text=cmd, size=11.5, font=MONO, color=MUTED)])
    if i < 2:
        rect(s, x + vw + gap / 2 - 0.14, vy + vh / 2 - 0.15, 0.28, 0.3, fill=GREY, kind=MSO_SHAPE.RIGHT_ARROW)

rect(s, M, 5.0, CW, 0.45, fill=SOFT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
textbox(s, M, 5.0, CW, 0.45, [dict(runs=[
    dict(text="THE LOOP     ", size=10, bold=True, color=MUTED, spc=100),
    dict(text="RECALL", bold=True, color=L3), dict(text=" at start  →  code with the agent  →  "),
    dict(text="CAPTURE", bold=True, color=L0), dict(text=" before resets  →  "),
    dict(text="CURATE", bold=True, color=L2), dict(text=" at task end  →  next task starts smarter"),
], size=12, color=INK, align=CENTER)], anchor=MID)

hw, hy, hh = (CW - 0.3) / 2, 5.6, 1.28
lx, rx = M, M + hw + 0.3
for x, c in ((lx, L3), (rx, L0)):
    rect(s, x, hy, hw, hh, fill=SOFT)
    rect(s, x, hy, 0.07, hh, fill=c)

textbox(s, lx + 0.2, hy + 0.08, hw - 0.4, 0.32, [dict(runs=[
    dict(text="Adopt in ~1 minute", size=13, bold=True, color=NAVY),
    dict(text="   ·   no build, no dependencies", size=11, color=MUTED),
])])
code = dict(font=MONO, size=10.5, color=NAVY)
textbox(s, lx + 0.2, hy + 0.42, hw - 0.4, hh - 0.45, [
    dict(runs=[dict(text="1   ", bold=True, color=L3), dict(text="Copy "), dict(text="memory-kit/", **code),
               dict(text=" into the repo (root or "), dict(text=".github/", **code), dict(text=")")], size=11, after=2),
    dict(runs=[dict(text="2   ", bold=True, color=L3), dict(text="Add one line to "), dict(text="AGENTS.md", **code),
               dict(text=", "), dict(text="CLAUDE.md", **code), dict(text=", "),
               dict(text="copilot-instructions.md", **code), dict(text=" or "), dict(text=".cursorrules", **code),
               dict(text=": “Follow the memory convention in memory-kit/MEMORY.md.”", italic=True)], size=11, after=2),
    dict(runs=[dict(text="3   ", bold=True, color=L3), dict(text="Set "), dict(text="<notes-dir>", **code),
               dict(text=" and "), dict(text="<docs-dir>", **code), dict(text="; git-ignore "),
               dict(text="notes/L0_*.md", **code)], size=11),
])

textbox(s, rx + 0.2, hy + 0.08, hw - 0.4, 0.32, [dict(text="Built-in guardrails", size=13, bold=True, color=NAVY)])
textbox(s, rx + 0.2, hy + 0.42, hw - 0.4, hh - 0.45, [dict(text=t, size=11, bullet=L0, after=1) for t in (
    "RECALL is read-only  ·  CAPTURE is capture-only  ·  CURATE is approval-first",
    "Lower-authority memory never overwrites higher-authority memory",
    "Anti-patterns are relocated, never silently deleted",
    "Missing sources are declared — coverage is never fabricated",
)])

os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT)
print(OUT)
